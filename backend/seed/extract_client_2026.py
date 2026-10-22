"""Extract the client's real 2026 project pipeline (ADR 0024) from their Excel
workbook into `client_2026_seed_data.json`, the default dataset of
`seed_demo_data.py`.

The workbook (`24-25-26 Global RPD Project Pipeline.xlsx`) is client data and
is never committed; only the JSON this script writes is. Treat the workbook as
untrusted input: run this with an isolated interpreter and pass the path as an
argument, from a directory that is not the workbook's own:

    uv run --no-project --with openpyxl python -I backend/seed/extract_client_2026.py \\
        /path/to/pipeline.xlsx [backend/seed/client_2026_seed_data.json] \\
        [--formula-workbook "reference/Formula-2025-26 Global RPD Project Pipeline.xlsx"]

The second workbook is the client's formula workbook (ADR 0035). It is
committed under `reference/` and is the default. Only its engineer roster
(`RPD Capacity!I39:T50`) is read.

It depends on the standard library plus `openpyxl` only (no backend imports),
so it runs under `python -I` without the backend's virtualenv.

Source and selection
--------------------
- Sheet `2024-25-26 Projects`. The header is row 5 and data starts at row 6.
  Only rows with a Project Name (column I) are read.
- ADR 0041 §1 (decision (a) of docs/DATA_RECONCILIATION_2026.md): the
  dataset is every row whose `Project Request (MM/YY)` (column B) falls in
  **2026**, **or** whose `Project to be considered for Capacity Calculation`
  (column D) is `Y`: 330 projects (253 requested in 2026 + 77 flagged rows
  requested in 2024/25, the carry-over). Every one of them is kept, including
  Cancelled, Commercialized and On Hold rows. Those statuses are displayed
  but never scheduled (DOMAIN_RULES "Booking rules"); every status but
  Cancelled / Draft counts in the full-year demand (I25).

Column mapping (0-based column index → JSON field)
--------------------------------------------------
- 0 `S.No.` → `s_no`. It is not unique in the sheet (596 repeats), so it is
  traceability only.
- 1 request date → `request_date` (ISO date) and `request_week` (ISO week).
  A `datetime` cell is used as is. A text cell `MM/YY` or `MM/YYYY` is parsed
  to the first of that month.
- 2 `Marketing Priority` (rank 1..13 or blank) → `marketing_priority`.
- 6 `Online PD Database Project Code`, 7 `RM Saving Project Codes` and 0 →
  `external_code`, which must be unique (`projects.external_code`). The first
  usable value wins:
  1. column 6 if it looks like `NN-NNNNN` (the PD database form);
  2. else column 7 unless it is one of `N/A`, `NA`, `Missing`, `-`, `TBD`,
     `TBA`, `TBR` or blank;
  3. else the generated code `26-S{S.No:04d}`.
  A code already taken by an earlier row (sheet order) falls through to the
  next rule. A generated code that is still taken gets `-2`, `-3`, ...
  `code_source` records which rule applied.
- 8 `Project Name` → `name`: whitespace collapsed, capped at 300 characters.
- 9 `Plant/OEM` → `plant`, 11 `Product Type` → `product_type`, 12
  `Requesting Department` → `requesting_department` (ADR 0035): the raw
  workbook value, whitespace collapsed, capped at 100 characters, null when
  blank. The values are kept as the workbook writes them (the controlled
  lists are `Lists & Explanations` columns C, Q and D,
  `domain_constants.PLANT_OPTIONS` etc.). They are not remapped: for
  example `OEM-SEITEK` stays the plant name as written while the hub is
  `OEM-Seitek`, and the department `Development`, which is not in the
  workbook's own list, is kept as written (docs/OPEN_QUESTIONS.md).
- 10 `RPDO hub` → `hub`. `OEM-Seitek` and `OEM-SEITEK` (the workbook's
  spellings) map to the app hub `OEM-Seitek` (ADR 0041 §4 renamed the app's
  `OEM-Seltek`). Every other value must be one of
  the six app hubs, otherwise the script fails loudly.
- 4 `Project Category` → `category`, with the OEM rule (ADR 0007): a project
  on an OEM hub carries `A-OEM`/`B-OEM`/`C-OEM` (`A+`/`A` → `A-OEM`, `B` →
  `B-OEM`, `C` → `C-OEM`). An `X-OEM` category on a non-OEM hub becomes `X`.
- 5 `Project Type` → `type`. `CO`, `NO`, `NM`, `IMP`, `RC`, `AC`
  (Alternative Component) and `SER` (Service) map one to one (ADR 0041 §4).
  `BC` (Business Continuity) is the app's `RC`, "Regulatory / Business
  Compliance" (ADR 0038 amendment 2026-10-22, owner ruling), and so are the
  text variants "Business Continuity", "Business Continiuty" and "Business
  Continuation" (any case). `CER` and `TCO` stay null (OQ #58). The raw code
  is always kept in `excel_type`.
- 18 `Certification Testing Required` → `excel_certification` (raw) and,
  ADR 0041 §3, `certification_testing_required = (col 18 == "Yes") OR
  (lab_effort_weeks > 0)`: the client's lab pivot sums col 21 with no cert
  filter (docs/DATA_RECONCILIATION_2026.md W9).
- ADR 0041 §3 per-project figures (number, 2 dp; a text cell holding a
  number is read as that number; anything else is null and counted):
  16 `Process Design Man Week` → `design_effort_weeks`, 21 `Lab Weeks-
  Process Derived` → `lab_effort_weeks` (both null when blank **or 0**: a
  typed 0 is a blank, the lead-time table applies), 22 `Lab Week for Safety
  testing` → `safety_lab_weeks` (0 kept), 17 `Actual Design Man Week` →
  `actual_design_weeks`, 23 `Actual Testing Time` → `actual_testing_weeks`
  (null when blank or 0), 24 `Safety Certification Required` →
  `safety_certification_required` (`yes`, any case).
- 32 `Project Status` → `status`: `In Que` → In Queue, `Development` → In
  Development, `Industrialization` → Under Industrialization, and
  `Commercialized`, `On Hold`, `Cancelled` unchanged. Any other value fails
  loudly.
- 33 `Project Review Remarks` → `comments`: trimmed, null when blank, capped
  at 4000 characters.
- 34 `Project Leader` → `leader`: trimmed. For a shared value such as
  `Lokesh/Rahul`, only the first name is kept. Null when blank. The leader
  roster is derived from this column (`leaders`, below).
- 44 `Customer` → `customer` (`ALL`/`COKE`/`KA`/`WM`, upper-cased; null when
  blank). This is the encrypted `customer_name` field. It is written here
  because it is the client's own data. It is never logged.
- 45 `Project End Date - LATEST` → `end_date` (ISO date, null when blank)
  and `target_end_week`. A `datetime` cell is used as is; a number is an
  Excel date serial (days since 1899-12-30) and is converted. Some rows
  hold the serial because the cell is formatted as a number. A text cell
  `Mon'YY` (e.g. `Feb'26`) is the first of that month (ADR 0041 §4); `-` is
  null.
  `target_end_week` is the date's ISO week when the ISO week-year is the
  plan year 2026 (W1–W52; a 2026 ISO week 53 becomes 52), and 52 + the ISO
  week when it is 2027 and that is still inside the 78-week horizon. **A
  date in an earlier ISO year (2025) is not imported** (`target_end_week`
  null): week numbers start at 2026 W1, so a 2025 date has no week. The
  raw `end_date` stays in the JSON (docs/OPEN_QUESTIONS.md, ADR 0035).
- 56 `Spillover '25` → `spillover_2025`: true for `Y` / `Yes` (any case),
  else false. Display only since ADR 0041 §1 (it no longer sets
  `carry_over`; ADR 0035 §10 superseded).
- 1 → `request_year`. ADR 0041 §1: `carry_over = request_year < 2026` and
  `reg_year = request_year` (the seed reads both).
- 52 `Development Step` → `development_step`. Kept only when it is a workflow
  step id (`PDD-A`..`PDD-N` or `OEM-A`..`OEM-N`). `Not Started` and blank
  give null.
- 3 `Project to be considered for Capacity Calculation` → `capacity_flag`
  (true for `Y`), stored as `projects.capacity_flag` (decision (a)).

The workbook has no financial values for these rows (TCOGS, selling price,
gross margin), so none are emitted.

Priority (no prioritization dimensions in the workbook)
-------------------------------------------------------
- `priority` is P1 for Marketing Priority 1–5 and P2 for 6–13.
- An unranked row gets a band from its status:
  - Under Industrialization or In Development → P3;
  - In Queue → P4;
  - the non-scheduled statuses (Commercialized, On Hold, Cancelled) → P4.
- The 13 dimension scores are derived deterministically by `seed_demo_data`
  so that their band equals `priority` (see `seed/client_2026.py`).

Leaders → engineers
-------------------
`leaders` lists every distinct leader with the hub where they lead the most
rows of this dataset. Ties go to the alphabetically first hub.

Roster (formula workbook, ADR 0035)
-----------------------------------
`roster` is the engineer table of the formula workbook's `RPD Capacity`
sheet: four blocks side by side, `I:K` India (PD), `L:N` India (R&D),
`O:Q` Greece (R&D), `R:T` Romania (PD). Row 39 names the block, row 41 is
the header (`Engineer Name`, `Resource Count`, `Remarks`) and the people
start at row 42. Each entry is `name`, `hub`, `fte` (the Resource Count),
`remark` and the source `cell`. Rows with a 0 Resource Count (the
placeholders "New Hiring" and "New engineer-Contractual") are kept here and
skipped by the seed, since an engineer needs FTE >= 0.01.

The seed (`seed/client_2026.py`) builds the engineers from `roster`, maps
the workbook leaders onto it, and adds the leaders who are not on it as
provisional engineers. It never invents an engineer or an assignment.
"""

from __future__ import annotations

import datetime as _dt
import json
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any

SHEET = "2024-25-26 Projects"
HEADER_ROW = 5
FIRST_DATA_ROW = 6
REQUEST_YEAR = 2026

# 0-based column indexes (header row 5).
COL_SNO = 0
COL_REQUEST = 1
COL_MKT_PRIORITY = 2
COL_CAPACITY_FLAG = 3
COL_CATEGORY = 4
COL_TYPE = 5
COL_PD_CODE = 6
COL_RM_CODE = 7
COL_NAME = 8
COL_PLANT = 9
COL_HUB = 10
COL_DESIGN_EFFORT = 16
COL_ACTUAL_DESIGN = 17
COL_CERT = 18
COL_LAB_EFFORT = 21
COL_SAFETY_LAB = 22
COL_ACTUAL_TESTING = 23
COL_SAFETY_CERT = 24
COL_STATUS = 32
COL_REMARKS = 33
COL_LEADER = 34
COL_PRODUCT_TYPE = 11
COL_REQUESTING_DEPARTMENT = 12
COL_CUSTOMER = 44
COL_END_DATE = 45
COL_DEV_STEP = 52
COL_SPILLOVER_25 = 56
#: The last column read; shorter rows are padded with None up to it.
COL_LAST = COL_SPILLOVER_25

#: The plan year's ISO week-year and the scheduler horizon (weeks).
PLAN_YEAR = 2026
PLAN_WEEKS = 52
HORIZON_WEEKS = 78

#: Header text expected at each used column. A changed layout fails loudly
#: instead of silently reading the wrong column.
EXPECTED_HEADERS = {
    COL_SNO: "S.No.",
    COL_REQUEST: "Project Request (MM/YY)",
    COL_MKT_PRIORITY: "Marketing Priority",
    COL_CATEGORY: "Project Category",
    COL_TYPE: "Project Type",
    COL_PD_CODE: "Online PD Database Project Code",
    COL_NAME: "Project Name",
    COL_HUB: "RPDO hub",
    COL_CERT: "Certification Testing Required",
    COL_CAPACITY_FLAG: "Project to be considered for Capacity Calculation",
    COL_DESIGN_EFFORT: "Process Design Man Week",
    COL_ACTUAL_DESIGN: "Actual Design Man Week",
    COL_LAB_EFFORT: "Lab Weeks- Process Derived",
    COL_SAFETY_LAB: "Lab Week for Safety testing",
    COL_ACTUAL_TESTING: "Actual Testing Time",
    COL_SAFETY_CERT: "Safety Certification Required",
    COL_STATUS: "Project Status",
    COL_LEADER: "Project Leader",
    COL_PLANT: "Plant/OEM",
    COL_PRODUCT_TYPE: "Product Type",
    COL_REQUESTING_DEPARTMENT: "Requesting Department",
    COL_CUSTOMER: "Customer",
    COL_END_DATE: "Project End Date - LATEST",
    COL_DEV_STEP: "Development Step",
    COL_SPILLOVER_25: "Spillover '25",
}

#: Formula workbook roster (ADR 0035): sheet, block header row, column
#: header row, first and last person rows, and per block (first column,
#: expected block label, app hub).
ROSTER_SHEET = "RPD Capacity"
ROSTER_LABEL_ROW = 39
ROSTER_HEADER_ROW = 41
ROSTER_FIRST_ROW = 42
ROSTER_LAST_ROW = 50
ROSTER_BLOCKS: tuple[tuple[str, str, str], ...] = (
    ("I", "India (PD)", "PD-India"),
    ("L", "India (R&D)", "R&D-India"),
    ("O", "Greece (R&D)", "R&D-Greece"),
    ("R", "Romania (PD)", "PD-Romania"),
)
DEFAULT_FORMULA_WORKBOOK = (
    Path(__file__).resolve().parents[2]
    / "reference"
    / "Formula-2025-26 Global RPD Project Pipeline.xlsx"
)

HUB_MAP = {
    "R&D-Greece": "R&D-Greece",
    "R&D-India": "R&D-India",
    "PD-India": "PD-India",
    "PD-Romania": "PD-Romania",
    "OEM-HCK": "OEM-HCK",
    "OEM-Seitek": "OEM-Seitek",
    # The prototype's spelling (renamed by ADR 0041 §4) and the workbook's
    # upper-case variant both normalise to the client's `OEM-Seitek`.
    "OEM-Seltek": "OEM-Seitek",
    "OEM-SEITEK": "OEM-Seitek",
}
OEM_HUBS = frozenset({"OEM-HCK", "OEM-Seitek"})

STATUS_MAP = {
    "In Que": "In Queue",
    "In Queue": "In Queue",
    "Development": "In Development",
    "Industrialization": "Under Industrialization",
    "Commercialized": "Commercialized",
    "On Hold": "On Hold",
    "Cancelled": "Cancelled",
}
SCHEDULED_STATUSES = frozenset({"In Queue", "In Development", "Under Industrialization"})

#: Workbook type (upper-cased, one line) → app `ProjectType`. ADR 0038
#: amendment (2026-10-22): `BC` and its text variants are `RC`.
TYPE_MAP = {
    "CO": "CO",
    "NO": "NO",
    "NM": "NM",
    "IMP": "IMP",
    "RC": "RC",
    "BC": "RC",
    # ADR 0041 §4: the client's own types.
    "AC": "AC",
    "SER": "SER",
    "BUSINESS CONTINUITY": "RC",
    "BUSINESS CONTINIUTY": "RC",
    "BUSINESS CONTINUATION": "RC",
}

_TO_OEM = {"A+": "A-OEM", "A": "A-OEM", "B": "B-OEM", "C": "C-OEM"}
_FROM_OEM = {"A-OEM": "A", "B-OEM": "B", "C-OEM": "C"}
_ALL_CATEGORIES = frozenset({"A+", "A", "B", "C", "A-OEM", "B-OEM", "C-OEM"})

_PD_CODE_RE = re.compile(r"^\d{2}-\d{3,6}$")
_RM_CODE_RE = re.compile(r"^[A-Za-z0-9_\-]{3,60}$")
_STEP_ID_RE = re.compile(r"^(PDD|OEM)-[A-N]$")
_NO_CODE = frozenset({"", "N/A", "NA", "MISSING", "-", "TBD", "TBA", "TBR"})
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")

NAME_MAX = 300
COMMENTS_MAX = 4000
LEADER_MAX = 200
#: Plant / product type / requesting department (`projects` columns, 100).
LABEL_MAX = 100

#: Excel's day zero for date serials (the 1900 system, after the leap bug).
_EXCEL_EPOCH = _dt.date(1899, 12, 30)
#: Plausible serial range (1954..2119) so a stray small number is not a date.
_EXCEL_SERIAL_RANGE = (20000, 80000)


def _text(value: Any) -> str:
    """A cell as clean text: control characters removed, whitespace trimmed."""

    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return _CONTROL_RE.sub("", str(value)).strip()


def _one_line(value: Any) -> str:
    return " ".join(_text(value).split())


def _request_date(value: Any) -> _dt.date | None:
    if isinstance(value, _dt.datetime):
        return value.date()
    if isinstance(value, _dt.date):
        return value
    text = _text(value)
    match = re.fullmatch(r"(\d{1,2})/(\d{2}|\d{4})", text)
    if not match:
        return None
    month, year = int(match.group(1)), int(match.group(2))
    if year < 100:
        year += 2000
    if not 1 <= month <= 12:
        return None
    return _dt.date(year, month, 1)


_MONTHS = {
    m: i
    for i, m in enumerate(
        ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"),
        start=1,
    )
}
_MON_YY_RE = re.compile(r"^([A-Z][a-z]{2})'(\d{2})$")


def _end_date(value: Any) -> _dt.date | None:
    """`Project End Date - LATEST`: a date cell, an Excel date serial, or a
    `Mon'YY` text (the first of that month, ADR 0041 §4). `-` is null."""

    if isinstance(value, _dt.datetime):
        return value.date()
    if isinstance(value, _dt.date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        low, high = _EXCEL_SERIAL_RANGE
        if low <= value <= high and float(value).is_integer():
            return _EXCEL_EPOCH + _dt.timedelta(days=int(value))
        return None
    match = _MON_YY_RE.fullmatch(_text(value))
    if match and match.group(1) in _MONTHS:
        return _dt.date(2000 + int(match.group(2)), _MONTHS[match.group(1)], 1)
    return None


def figure(value: Any, *, keep_zero: bool = False) -> float | None:
    """A per-project figure (ADR 0041 §3): a number (or a text cell holding
    one) rounded to 2 dp; null when blank, non-numeric, negative, or 0 unless
    `keep_zero`. A typed 0 is a blank for the effort figures."""

    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = _text(value)
        if not text:
            return None
        try:
            number = float(text)
        except ValueError:
            return None
    if number != number or number < 0:  # NaN or negative
        return None
    number = round(number, 2)
    if number == 0 and not keep_zero:
        return None
    return number


def plan_week(date: _dt.date | None) -> int | None:
    """The plan week of a date (module docstring, column 45): ISO weeks of
    2026 are 1–52 (week 53 → 52), ISO weeks of 2027 follow as 53–78. Any
    other year, or past the horizon, gives None."""

    if date is None:
        return None
    iso_year, iso_week, _ = date.isocalendar()
    if iso_year == PLAN_YEAR:
        return min(iso_week, PLAN_WEEKS)
    if iso_year == PLAN_YEAR + 1 and PLAN_WEEKS + iso_week <= HORIZON_WEEKS:
        return PLAN_WEEKS + iso_week
    return None


def _label(value: Any) -> str | None:
    return _one_line(value)[:LABEL_MAX] or None


def _marketing_priority(value: Any) -> int | None:
    text = _text(value)
    if not text:
        return None
    try:
        rank = float(text)
    except ValueError:
        return None
    if not rank.is_integer() or rank < 1:
        return None
    return int(rank)


def priority_for(marketing_priority: int | None, status: str) -> str:
    """Band from the Marketing Priority rank, else from the status (docstring)."""

    if marketing_priority is not None:
        if marketing_priority <= 5:
            return "P1"
        if marketing_priority <= 13:
            return "P2"
    if status in ("Under Industrialization", "In Development"):
        return "P3"
    return "P4"


def category_for(hub: str, excel_category: str) -> str:
    """The OEM category rule (ADR 0007), both directions."""

    if excel_category not in _ALL_CATEGORIES:
        raise ValueError(f"unknown Project Category {excel_category!r}")
    if hub in OEM_HUBS:
        return _TO_OEM.get(excel_category, excel_category)
    return _FROM_OEM.get(excel_category, excel_category)


def leader_for(raw: Any) -> str | None:
    text = _one_line(raw)
    if not text:
        return None
    first = text.split("/")[0].strip()
    return first[:LEADER_MAX] or None


def _header_matches(cell: Any, expected: str) -> bool:
    return " ".join(_text(cell).split()).lower() == expected.lower()


def read_rows(xlsx_path: Path) -> list[tuple[int, tuple[Any, ...]]]:
    """(excel_row_number, values) for every data row with a Project Name."""

    import openpyxl  # imported here so the module's helpers need no openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb[SHEET]
        header = next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True))
        for col, expected in EXPECTED_HEADERS.items():
            if not _header_matches(header[col], expected):
                raise ValueError(f"column {col} header is {header[col]!r}, expected {expected!r}")
        out: list[tuple[int, tuple[Any, ...]]] = []
        for number, values in enumerate(
            ws.iter_rows(min_row=FIRST_DATA_ROW, max_col=COL_LAST + 1, values_only=True),
            start=FIRST_DATA_ROW,
        ):
            padded = tuple(values) + (None,) * max(0, COL_LAST + 1 - len(values))
            if _one_line(padded[COL_NAME]):
                out.append((number, padded))
        return out
    finally:
        wb.close()


def extract(rows: list[tuple[int, tuple[Any, ...]]]) -> dict[str, Any]:
    """Map the sheet rows to the seed JSON (module docstring). Pure."""

    projects: list[dict[str, Any]] = []
    used_codes: set[str] = set()
    for number, row in rows:
        request = _request_date(row[COL_REQUEST])
        capacity_flag = _one_line(row[COL_CAPACITY_FLAG]).upper() == "Y"
        # ADR 0041 §1: requested in 2026, or flagged for capacity (any year).
        in_2026 = request is not None and request.year == REQUEST_YEAR
        if not (in_2026 or (capacity_flag and request is not None)):
            continue
        assert request is not None

        name = _one_line(row[COL_NAME])[:NAME_MAX]
        excel_hub = _one_line(row[COL_HUB])
        if excel_hub not in HUB_MAP:
            raise ValueError(f"row {number}: unknown RPDO hub {excel_hub!r}")
        hub = HUB_MAP[excel_hub]

        excel_status = _one_line(row[COL_STATUS])
        if excel_status not in STATUS_MAP:
            raise ValueError(f"row {number}: unknown Project Status {excel_status!r}")
        status = STATUS_MAP[excel_status]

        excel_category = _one_line(row[COL_CATEGORY])
        category = category_for(hub, excel_category)

        excel_type = _one_line(row[COL_TYPE]).upper()
        marketing_priority = _marketing_priority(row[COL_MKT_PRIORITY])

        s_no = _text(row[COL_SNO])
        pd_code = _one_line(row[COL_PD_CODE])
        rm_code = _one_line(row[COL_RM_CODE])
        if _PD_CODE_RE.fullmatch(pd_code) and pd_code not in used_codes:
            code, source = pd_code, "pd_database"
        elif (
            rm_code.upper() not in _NO_CODE
            and _RM_CODE_RE.fullmatch(rm_code)
            and rm_code not in used_codes
        ):
            code, source = rm_code, "rm_saving"
        else:
            base = f"26-S{int(s_no):04d}" if s_no.isdigit() else f"26-R{number:04d}"
            code, n = base, 1
            while code in used_codes:
                n += 1
                code = f"{base}-{n}"
            source = "generated"
        used_codes.add(code)

        dev_step = _one_line(row[COL_DEV_STEP]).upper()
        customer = _one_line(row[COL_CUSTOMER]).upper()
        end_date = _end_date(row[COL_END_DATE])
        comments = _text(row[COL_REMARKS])[:COMMENTS_MAX]
        excel_cert = _one_line(row[COL_CERT])
        lab_effort = figure(row[COL_LAB_EFFORT])
        projects.append(
            {
                "excel_row": number,
                "s_no": s_no,
                "external_code": code,
                "code_source": source,
                "name": name,
                "hub": hub,
                "excel_hub": excel_hub,
                "category": category,
                "excel_category": excel_category,
                "type": TYPE_MAP.get(excel_type),
                "excel_type": excel_type or None,
                "status": status,
                "excel_status": excel_status,
                "marketing_priority": marketing_priority,
                "priority": priority_for(marketing_priority, status),
                "excel_certification": excel_cert or None,
                # ADR 0041 §3: cert Yes, or the project carries lab weeks.
                "certification_testing_required": excel_cert.lower() == "yes"
                or (lab_effort is not None and lab_effort > 0),
                "design_effort_weeks": figure(row[COL_DESIGN_EFFORT]),
                "lab_effort_weeks": lab_effort,
                "safety_lab_weeks": figure(row[COL_SAFETY_LAB], keep_zero=True),
                "actual_design_weeks": figure(row[COL_ACTUAL_DESIGN]),
                "actual_testing_weeks": figure(row[COL_ACTUAL_TESTING]),
                "safety_certification_required": _one_line(row[COL_SAFETY_CERT]).lower() == "yes",
                "leader": leader_for(row[COL_LEADER]),
                "customer": customer or None,
                "comments": comments or None,
                "development_step": dev_step if _STEP_ID_RE.fullmatch(dev_step) else None,
                "request_date": request.isoformat(),
                "request_week": request.isocalendar()[1],
                "plant": _label(row[COL_PLANT]),
                "product_type": _label(row[COL_PRODUCT_TYPE]),
                "requesting_department": _label(row[COL_REQUESTING_DEPARTMENT]),
                "end_date": end_date.isoformat() if end_date is not None else None,
                "target_end_week": plan_week(end_date),
                "spillover_2025": _one_line(row[COL_SPILLOVER_25]).upper() in ("Y", "YES"),
                "capacity_flag": capacity_flag,
                "request_year": request.year,
            }
        )

    hub_counts: dict[str, Counter[str]] = {}
    for p in projects:
        if p["leader"]:
            hub_counts.setdefault(p["leader"], Counter())[p["hub"]] += 1
    leaders = []
    for name in sorted(hub_counts):
        counts = hub_counts[name]
        home = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))[0][0]
        leaders.append(
            {
                "name": name,
                "hub": home,
                "project_count": sum(counts.values()),
                "hub_counts": dict(sorted(counts.items())),
            }
        )

    return {
        "source": {
            "workbook": "24-25-26 Global RPD Project Pipeline.xlsx (client; not committed)",
            "sheet": SHEET,
            "selection": (
                f"Project Request (MM/YY) in {REQUEST_YEAR}, or Project to be considered "
                "for Capacity Calculation = Y (ADR 0041 §1)"
            ),
            "extractor": "backend/seed/extract_client_2026.py",
        },
        "project_count": len(projects),
        "leaders": leaders,
        "projects": projects,
    }


def _column_index(letter: str) -> int:
    """1-based index of a single-letter column (A..Z)."""

    return ord(letter.upper()) - ord("A") + 1


def read_roster(xlsx_path: Path) -> list[dict[str, Any]]:
    """The formula workbook's engineer roster (module docstring "Roster")."""

    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb[ROSTER_SHEET]
        rows = {
            number: tuple(values)
            for number, values in enumerate(
                ws.iter_rows(
                    min_row=ROSTER_LABEL_ROW,
                    max_row=ROSTER_LAST_ROW,
                    max_col=_column_index("T"),
                    values_only=True,
                ),
                start=ROSTER_LABEL_ROW,
            )
        }
    finally:
        wb.close()
    return extract_roster(rows)


def extract_roster(rows: dict[int, tuple[Any, ...]]) -> list[dict[str, Any]]:
    """Map the `RPD Capacity` rows 39..50 (1-based row → values from column
    A) to roster entries. Pure; fails loudly on a changed layout."""

    def cell(row: int, col: int) -> Any:
        values = rows.get(row, ())
        return values[col - 1] if col - 1 < len(values) else None

    roster: list[dict[str, Any]] = []
    for letter, label, hub in ROSTER_BLOCKS:
        col = _column_index(letter)
        if not _header_matches(cell(ROSTER_LABEL_ROW, col), label):
            raise ValueError(f"roster block {letter}{ROSTER_LABEL_ROW} is not {label!r}")
        for offset, expected in enumerate(("Engineer Name", "Resource Count", "Remarks")):
            if not _header_matches(cell(ROSTER_HEADER_ROW, col + offset), expected):
                raise ValueError(f"roster header at column {col + offset} is not {expected!r}")
        for row in range(ROSTER_FIRST_ROW, ROSTER_LAST_ROW + 1):
            name = _one_line(cell(row, col))
            if not name:
                continue
            raw_fte = cell(row, col + 1)
            if not isinstance(raw_fte, (int, float)) or isinstance(raw_fte, bool):
                raise ValueError(f"roster {letter}{row}: Resource Count is not a number")
            roster.append(
                {
                    "name": name[:LEADER_MAX],
                    "hub": hub,
                    "fte": round(float(raw_fte), 2),
                    "remark": _one_line(cell(row, col + 2)) or None,
                    "cell": f"{letter}{row}:{chr(ord(letter) + 1)}{row}",
                }
            )
    return roster


def main(argv: list[str]) -> int:
    import argparse

    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("pipeline", type=Path, help="24-25-26 Global RPD Project Pipeline.xlsx")
    parser.add_argument(
        "out",
        type=Path,
        nargs="?",
        default=Path(__file__).resolve().parent / "client_2026_seed_data.json",
    )
    parser.add_argument("--formula-workbook", type=Path, default=DEFAULT_FORMULA_WORKBOOK)
    args = parser.parse_args(argv[1:])
    data = extract(read_rows(args.pipeline))
    data["source"]["roster"] = (
        "Formula-2025-26 Global RPD Project Pipeline.xlsx (client; reference/), "
        f"{ROSTER_SHEET}!I{ROSTER_LABEL_ROW}:T{ROSTER_LAST_ROW}"
    )
    data["roster"] = read_roster(args.formula_workbook)
    args.out.write_text(json.dumps(data, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    # Counts only: never print row values (customer is an encrypted field).
    projects = data["projects"]
    print(
        f"wrote {data['project_count']} projects, {len(data['leaders'])} leaders, "
        f"{len(data['roster'])} roster rows to {args.out}"
    )
    # ADR 0041 §7.1 summary: counts only (never a row value).
    by_hub_status = Counter((p["hub"], p["status"]) for p in projects)
    cert_by_lab = sum(
        1
        for p in projects
        if p["certification_testing_required"]
        and (p["excel_certification"] or "").lower() != "yes"
    )
    print(
        "carry-over (request year < 2026): "
        f"{sum(1 for p in projects if p['request_year'] < REQUEST_YEAR)}; "
        f"capacity flag Y: {sum(1 for p in projects if p['capacity_flag'])}; "
        f"blank design effort: {sum(1 for p in projects if p['design_effort_weeks'] is None)}; "
        f"blank lab effort: {sum(1 for p in projects if p['lab_effort_weeks'] is None)}; "
        f"cert set by lab weeks only: {cert_by_lab}"
    )
    for (hub, status), n in sorted(by_hub_status.items()):
        print(f"  {hub:12s} {status:24s} {n}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
