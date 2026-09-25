"""P5-T04 — `GET /exports/{surface}` (sync, streamed) and the async
(`?async_export=true`) job dispatch/status/download endpoints
(`api/routers/exports.py`).

Same conventions as `tests/test_scenario_apply_api.py`/
`tests/test_hub_scope.py`: a real DB-backed `User` (for
`AuditLogEntry.actor_user_id`'s FK and `ExportJob.requested_by_user_id`'s
FK) plus an overridden `Principal` via `tests.auth_helpers`.

**What this file does NOT cover** — the async path's actual file
generation/MinIO upload (`workers.export_tasks.generate_export`'s task
body): `generate_export.delay(...)` is monkeypatched here to a fake that
just records its call args, mirroring `tests/
test_schedule_runs_cp_sat_api.py`'s established "never spin up a real
Celery worker" precedent for the exact same reason (Celery's real broker/
worker round-trip is not meaningfully testable against the SAVEPOINT-scoped
`db_session` fixture — a real worker process would open its own, separate
DB connection that can never see this fixture's uncommitted transaction).
The task body itself, including a REAL MinIO upload/download round-trip, is
covered separately in `tests/test_export_worker_task.py`, which spins up
its own dedicated (real, committed) Postgres + MinIO containers.
"""

from __future__ import annotations

import csv
import io
import uuid

import openpyxl
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select

from api.db import get_db
from api.main import app
from models.audit import AuditLogEntry
from models.enums import (
    ExportFormat,
    ExportJobStatus,
    ExportSurface,
    HubName,
    LabRegion,
    RoleName,
)
from models.export import ExportJob
from models.schedule import ScheduleRunProjectStep
from tests.auth_helpers import (
    clear_current_principal_override,
    make_principal,
    override_current_principal,
)
from tests.factories import (
    make_engineer,
    make_hub,
    make_priority_score,
    make_project,
    make_schedule_run,
    make_user,
    make_workflow_step_template,
)


def _client(db_session) -> AsyncClient:
    async def _override_get_db():
        yield db_session

    app.dependency_overrides[get_db] = _override_get_db
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


def _teardown():
    app.dependency_overrides.pop(get_db, None)
    clear_current_principal_override()


async def _as(db_session, *roles: RoleName, hub_scope_all: bool = True, hub_ids=()):
    user = await make_user(db_session, *roles, hub_scope_all=hub_scope_all, hub_ids=list(hub_ids))
    client = _client(db_session)
    override_current_principal(
        make_principal(*roles, user_id=user.id, hub_scope_all=hub_scope_all, hub_ids=hub_ids)
    )
    return user, client


def _parse_csv(content: bytes) -> list[dict]:
    text = content.decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text)))


def _parse_xlsx_sheet(content: bytes, sheet_name: str) -> tuple[list[str], list[list]]:
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    ws = workbook[sheet_name]
    all_rows = list(ws.iter_rows(values_only=True))
    return list(all_rows[0]), [list(r) for r in all_rows[1:]]


# --------------------------------------------------------------------------
# Synchronous CSV/XLSX exports — RBAC + hub scoping + data shape
# --------------------------------------------------------------------------


async def test_dashboard_export_csv_hub_scoped(db_session):
    hub_a = await make_hub(db_session, name=HubName.RD_GREECE, lab_region=LabRegion.GREECE)
    hub_b = await make_hub(db_session, name=HubName.PD_ROMANIA, lab_region=LabRegion.ROMANIA)
    await make_project(db_session, hub_a, name="Dash A")
    await make_project(db_session, hub_b, name="Dash B")

    try:
        _, client = await _as(
            db_session, RoleName.HUB_PLANNER, hub_scope_all=False, hub_ids=[hub_a.id]
        )
        resp = await client.get("/exports/dashboard", params={"format": "csv"})
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("text/csv")
        assert 'filename="dashboard-export.csv"' in resp.headers["content-disposition"]
        rows = _parse_csv(resp.content)
        names = {r["project_name"] for r in rows}
        assert names == {"Dash A"}
    finally:
        _teardown()


async def test_dashboard_export_requires_dashboard_read_permission(db_session):
    await make_hub(db_session)
    try:
        _, client = await _as(db_session, RoleName.ENGINEER, hub_scope_all=False, hub_ids=[])
        resp = await client.get("/exports/dashboard")
        assert resp.status_code == 403
    finally:
        _teardown()


async def test_capacity_export_csv_no_active_run_still_200s(db_session):
    await make_hub(db_session, name=HubName.RD_GREECE, lab_region=LabRegion.GREECE)
    try:
        _, client = await _as(db_session, RoleName.PORTFOLIO_MANAGER)
        resp = await client.get("/exports/capacity", params={"format": "csv"})
        assert resp.status_code == 200
        rows = _parse_csv(resp.content)
        assert len(rows) == 1  # exactly the one hub created above
    finally:
        _teardown()


async def test_matrix_export_xlsx_includes_financial_columns(db_session):
    hub = await make_hub(db_session, name=HubName.RD_GREECE, lab_region=LabRegion.GREECE)
    project = await make_project(
        db_session,
        hub,
        name="Matrix Project",
        tcogs_eur=1234.5,
        selling_price_eur=2000.0,
        gross_margin_pct=35.5,
        customer_name="Acme Co",
    )
    await make_priority_score(db_session, project, dims=[4] * 13)

    try:
        _, client = await _as(db_session, RoleName.PORTFOLIO_MANAGER)
        resp = await client.get("/exports/matrix", params={"format": "xlsx"})
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument"
        )
        headers, rows = _parse_xlsx_sheet(resp.content, "Matrix")
        assert "tcogs_eur" in headers
        assert "selling_price_eur" in headers
        assert "gross_margin_pct" in headers
        row = dict(zip(headers, rows[0], strict=True))
        assert row["project_name"] == "Matrix Project"
        assert float(row["tcogs_eur"]) == pytest.approx(1234.5)
        assert float(row["gross_margin_pct"]) == pytest.approx(35.5)
    finally:
        _teardown()


async def test_matrix_export_requires_matrix_read_not_write(db_session):
    """Hub Planner has MATRIX/READ (not WRITE, per the RBAC matrix) —
    confirms the export uses READ, not the stricter write permission.
    """
    hub = await make_hub(db_session, name=HubName.RD_GREECE, lab_region=LabRegion.GREECE)
    await make_project(db_session, hub, name="P")
    try:
        _, client = await _as(
            db_session, RoleName.HUB_PLANNER, hub_scope_all=False, hub_ids=[hub.id]
        )
        resp = await client.get("/exports/matrix", params={"format": "csv"})
        assert resp.status_code == 200
    finally:
        _teardown()


async def test_gantt_export_flattens_project_and_step_rows(db_session):
    hub = await make_hub(db_session, name=HubName.RD_GREECE, lab_region=LabRegion.GREECE)
    template_a = await make_workflow_step_template(db_session, step_id="PDD-A", sequence_order=1)
    template_b = await make_workflow_step_template(
        db_session, step_id="PDD-B", name="Concept Study", sequence_order=2
    )
    engineer = await make_engineer(db_session, hub)
    project = await make_project(db_session, hub, name="Gantt Project", leader=engineer)
    run = await make_schedule_run(db_session, is_active=True)
    db_session.add_all(
        [
            ScheduleRunProjectStep(
                schedule_run_id=run.id,
                project_id=project.id,
                step_template_id=template_a.id,
                sequence_order=1,
                duration_weeks=2,
                start_week=31,
                end_week=32,
                assigned_engineer_id=engineer.id,
            ),
            ScheduleRunProjectStep(
                schedule_run_id=run.id,
                project_id=project.id,
                step_template_id=template_b.id,
                sequence_order=2,
                duration_weeks=2,
                start_week=33,
                end_week=34,
                assigned_engineer_id=engineer.id,
            ),
        ]
    )
    await db_session.flush()

    try:
        _, client = await _as(db_session, RoleName.PORTFOLIO_MANAGER)
        resp = await client.get("/exports/gantt", params={"format": "csv"})
        assert resp.status_code == 200
        rows = _parse_csv(resp.content)
        assert len(rows) == 2  # one row per step, project fields repeated
        assert {r["step_id"] for r in rows} == {"PDD-A", "PDD-B"}
        assert all(r["project_name"] == "Gantt Project" for r in rows)
    finally:
        _teardown()


async def test_project_registration_export_includes_financial_fields(db_session):
    hub = await make_hub(db_session, name=HubName.RD_GREECE, lab_region=LabRegion.GREECE)
    await make_project(
        db_session,
        hub,
        name="Reg Project",
        customer_name="Beta Customer",
        tcogs_eur=500.25,
    )
    try:
        _, client = await _as(db_session, RoleName.PORTFOLIO_MANAGER)
        resp = await client.get("/exports/project-registration", params={"format": "csv"})
        assert resp.status_code == 200
        rows = _parse_csv(resp.content)
        assert rows[0]["customer_name"] == "Beta Customer"
        assert float(rows[0]["tcogs_eur"]) == pytest.approx(500.25)
    finally:
        _teardown()


async def test_capacity_planning_export_csv_entity_selector(db_session):
    hub = await make_hub(db_session, name=HubName.RD_GREECE, lab_region=LabRegion.GREECE)
    await make_engineer(db_session, hub, name="Eng One")

    try:
        _, client = await _as(db_session, RoleName.PORTFOLIO_MANAGER)

        resp_eng = await client.get(
            "/exports/capacity-planning", params={"format": "csv", "entity": "engineers"}
        )
        assert resp_eng.status_code == 200
        eng_rows = _parse_csv(resp_eng.content)
        assert eng_rows[0]["name"] == "Eng One"

        resp_ch = await client.get(
            "/exports/capacity-planning", params={"format": "csv", "entity": "chambers"}
        )
        assert resp_ch.status_code == 200
        ch_rows = _parse_csv(resp_ch.content)
        assert ch_rows == []  # no chambers created in this test
    finally:
        _teardown()


async def test_capacity_planning_export_xlsx_has_both_sheets(db_session):
    hub = await make_hub(db_session, name=HubName.RD_GREECE, lab_region=LabRegion.GREECE)
    await make_engineer(db_session, hub, name="Eng Two")

    try:
        _, client = await _as(db_session, RoleName.PORTFOLIO_MANAGER)
        resp = await client.get("/exports/capacity-planning", params={"format": "xlsx"})
        assert resp.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(resp.content), read_only=True)
        assert set(workbook.sheetnames) == {"Engineers", "Chambers"}
    finally:
        _teardown()


# --------------------------------------------------------------------------
# Async dispatch (?async_export=true) — job creation, audit log, RBAC reuse
# --------------------------------------------------------------------------


async def test_async_export_dispatch_creates_queued_job_and_audit_row(db_session, monkeypatch):
    hub = await make_hub(db_session, name=HubName.RD_GREECE, lab_region=LabRegion.GREECE)
    await make_project(db_session, hub, name="Async Project")

    captured = {}

    class _FakeAsyncResult:
        id = "fake-task-id-123"

    def _fake_delay(job_id, **kwargs):
        captured["job_id"] = job_id
        captured["kwargs"] = kwargs
        return _FakeAsyncResult()

    import workers.export_tasks as export_tasks_module

    monkeypatch.setattr(export_tasks_module.generate_export, "delay", _fake_delay)

    try:
        user, client = await _as(
            db_session, RoleName.HUB_PLANNER, hub_scope_all=False, hub_ids=[hub.id]
        )
        resp = await client.get(
            "/exports/dashboard", params={"format": "xlsx", "async_export": "true"}
        )
        assert resp.status_code == 202
        body = resp.json()
        assert body["status"] == "queued"
        job_id = uuid.UUID(body["job_id"])

        assert captured["job_id"] == str(job_id)
        assert captured["kwargs"]["surface"] == "dashboard"
        assert captured["kwargs"]["format"] == "xlsx"
        assert captured["kwargs"]["principal"]["user_id"] == str(user.id)
        assert captured["kwargs"]["principal"]["hub_ids"] == [str(hub.id)]

        job = await db_session.get(ExportJob, job_id)
        assert job is not None
        assert job.status.value == "queued"
        assert job.requested_by_user_id == user.id
        assert job.celery_task_id == "fake-task-id-123"
        assert job.filters["hub_id"] is None  # no explicit hub_id query param was sent

        audit_row = (
            await db_session.execute(
                select(AuditLogEntry).where(AuditLogEntry.entity_id == str(job_id))
            )
        ).scalar_one_or_none()
        assert audit_row is not None
        assert audit_row.action == "export.job_dispatch"
        assert audit_row.after_state["surface"] == "dashboard"
    finally:
        _teardown()


async def test_async_export_dispatch_requires_permission(db_session, monkeypatch):
    import workers.export_tasks as export_tasks_module

    monkeypatch.setattr(
        export_tasks_module.generate_export,
        "delay",
        lambda *a, **k: pytest.fail("must not dispatch when RBAC denies"),
    )
    try:
        _, client = await _as(db_session, RoleName.ENGINEER, hub_scope_all=False, hub_ids=[])
        resp = await client.get("/exports/dashboard", params={"async_export": "true"})
        assert resp.status_code == 403
    finally:
        _teardown()


# --------------------------------------------------------------------------
# Job status / download — owner-scoped, not surface-RBAC-scoped
# --------------------------------------------------------------------------


async def test_export_job_status_owner_can_read_own_job(db_session):
    user = await make_user(db_session, RoleName.PORTFOLIO_MANAGER)
    job = ExportJob(
        surface=ExportSurface.DASHBOARD,
        format=ExportFormat.CSV,
        requested_by_user_id=user.id,
        filters={},
        status=ExportJobStatus.QUEUED,
    )
    db_session.add(job)
    await db_session.flush()

    try:
        client = _client(db_session)
        override_current_principal(
            make_principal(RoleName.PORTFOLIO_MANAGER, user_id=user.id)
        )
        resp = await client.get(f"/exports/jobs/{job.id}")
        assert resp.status_code == 200
        body = resp.json()
        assert body["status"] == "queued"
        assert body["surface"] == "dashboard"
    finally:
        _teardown()


async def test_export_job_status_404_for_non_owner_non_admin(db_session):
    owner = await make_user(db_session, RoleName.PORTFOLIO_MANAGER, email="owner@example.com")
    other = await make_user(db_session, RoleName.PORTFOLIO_MANAGER, email="other@example.com")
    job = ExportJob(
        surface=ExportSurface.DASHBOARD,
        format=ExportFormat.CSV,
        requested_by_user_id=owner.id,
        filters={},
        status=ExportJobStatus.QUEUED,
    )
    db_session.add(job)
    await db_session.flush()

    try:
        client = _client(db_session)
        override_current_principal(
            make_principal(RoleName.PORTFOLIO_MANAGER, user_id=other.id)
        )
        resp = await client.get(f"/exports/jobs/{job.id}")
        assert resp.status_code == 404
    finally:
        _teardown()


async def test_export_job_status_admin_can_read_others_job(db_session):
    owner = await make_user(db_session, RoleName.PORTFOLIO_MANAGER, email="owner2@example.com")
    admin = await make_user(db_session, RoleName.ADMIN, email="admin@example.com")
    job = ExportJob(
        surface=ExportSurface.MATRIX,
        format=ExportFormat.XLSX,
        requested_by_user_id=owner.id,
        filters={},
        status=ExportJobStatus.COMPLETED,
        row_count=5,
    )
    db_session.add(job)
    await db_session.flush()

    try:
        client = _client(db_session)
        override_current_principal(make_principal(RoleName.ADMIN, user_id=admin.id))
        resp = await client.get(f"/exports/jobs/{job.id}")
        assert resp.status_code == 200
        assert resp.json()["row_count"] == 5
    finally:
        _teardown()


async def test_export_job_status_unknown_id_404s(db_session):
    user = await make_user(db_session, RoleName.PORTFOLIO_MANAGER)
    try:
        client = _client(db_session)
        override_current_principal(make_principal(RoleName.PORTFOLIO_MANAGER, user_id=user.id))
        resp = await client.get(f"/exports/jobs/{uuid.uuid4()}")
        assert resp.status_code == 404
    finally:
        _teardown()


async def test_export_download_409_when_job_not_yet_completed(db_session):
    user = await make_user(db_session, RoleName.PORTFOLIO_MANAGER)
    job = ExportJob(
        surface=ExportSurface.DASHBOARD,
        format=ExportFormat.CSV,
        requested_by_user_id=user.id,
        filters={},
        status=ExportJobStatus.RUNNING,
    )
    db_session.add(job)
    await db_session.flush()

    try:
        client = _client(db_session)
        override_current_principal(make_principal(RoleName.PORTFOLIO_MANAGER, user_id=user.id))
        resp = await client.get(f"/exports/jobs/{job.id}/download")
        assert resp.status_code == 409
    finally:
        _teardown()


async def test_export_download_404_for_non_owner(db_session):
    owner = await make_user(db_session, RoleName.PORTFOLIO_MANAGER, email="owner3@example.com")
    other = await make_user(db_session, RoleName.PORTFOLIO_MANAGER, email="other3@example.com")
    job = ExportJob(
        surface=ExportSurface.DASHBOARD,
        format=ExportFormat.CSV,
        requested_by_user_id=owner.id,
        filters={},
        status=ExportJobStatus.COMPLETED,
        object_key="exports/whatever.csv",
    )
    db_session.add(job)
    await db_session.flush()

    try:
        client = _client(db_session)
        override_current_principal(make_principal(RoleName.PORTFOLIO_MANAGER, user_id=other.id))
        resp = await client.get(f"/exports/jobs/{job.id}/download")
        assert resp.status_code == 404
    finally:
        _teardown()
