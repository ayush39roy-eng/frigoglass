"""P5-T04 — `workers.export_tasks._generate_export_async` (the async
"large export, via MinIO" task body), exercised against REAL, dedicated
Postgres 17 + MinIO containers — not the SAVEPOINT-scoped shared
`db_session` fixture (`tests/conftest.py`) and not a mocked MinIO client.

Mirrors `tests/test_schedule_runs_progress_sse.py` / `tests/
test_workers_progress.py`'s established "own dedicated throwaway
container(s), settings/client reset via `.cache_clear()` + poking the
module-level singleton back to `None`" pattern, and `tests/
test_seed_scripts.py`'s "own dedicated Postgres container, real commits, not
a per-test SAVEPOINT" pattern — this task genuinely needs both: a real,
committed DB row this task's OWN fresh `AsyncEngine` (opened inside
`_generate_export_async` itself, on a brand-new event loop each call, per
that module's docstring) can see, and a real object-storage round-trip
(upload from the task, download via `api.routers.exports.download_export_job`
called directly) — neither is meaningfully mockable without weakening
exactly the two things this test exists to prove.

`tests/test_export_api.py` deliberately does NOT cover this — see that
file's own module docstring for why (`generate_export.delay(...)` is
monkeypatched there instead).
"""

from __future__ import annotations

import io
import uuid
from types import SimpleNamespace

import openpyxl
import pytest
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from testcontainers.core.container import DockerContainer
from testcontainers.core.waiting_utils import wait_for_logs
from testcontainers.postgres import PostgresContainer

import core.minio_config as minio_config_module
from api.routers.exports import download_export_job
from core.minio_config import get_minio_client, get_minio_settings
from core.principal import Principal
from models.enums import (
    ExportFormat,
    ExportJobStatus,
    ExportSurface,
    HubName,
    LabRegion,
    ProjectCategory,
    ProjectPriority,
    ProjectStatus,
    ProjectType,
    RoleName,
)
from models.export import ExportJob
from models.hub import Hub
from models.priority import DIMENSION_FIELD_NAMES, PriorityScore
from models.project import Project
from models.user import User
from tests.conftest import run_alembic
from workers.export_tasks import _generate_export_async

_MINIO_ACCESS_KEY = "rpdtestaccess"
_MINIO_SECRET_KEY = "rpdtestsecret123"


class _FakeTask:
    """The only attribute `_generate_export_async` reads off its `task`
    argument (`task.request.id`) — a real `celery.Task` is not needed since
    this test calls the async body directly, bypassing Celery's broker
    entirely (mirroring `api/routers/schedule_runs.py`'s own
    `_progress_event_source` direct-call precedent, see that file's
    docstring point 6 / `tests/test_schedule_runs_progress_sse.py`).
    """

    request = SimpleNamespace(id="fake-worker-task-id")


@pytest.fixture(scope="module")
def postgres_url():
    with PostgresContainer("postgres:17", driver="asyncpg") as pg:
        url = pg.get_connection_url()
        r = run_alembic(["upgrade", "head"], url)
        assert r.returncode == 0, f"alembic upgrade head failed:\n{r.stdout}\n{r.stderr}"
        yield url


@pytest.fixture(scope="module")
def minio_endpoint():
    container = (
        DockerContainer("minio/minio")
        .with_exposed_ports(9000)
        .with_env("MINIO_ROOT_USER", _MINIO_ACCESS_KEY)
        .with_env("MINIO_ROOT_PASSWORD", _MINIO_SECRET_KEY)
        .with_command("server /data")
    )
    container.start()
    try:
        wait_for_logs(container, "API:", timeout=30)
        host = container.get_container_host_ip()
        port = container.get_exposed_port(9000)
        yield f"{host}:{port}"
    finally:
        container.stop()


@pytest.fixture
def _point_at_real_infra(postgres_url, minio_endpoint, monkeypatch):
    """Overrides `RPD_DATABASE_URL`/`RPD_MINIO_*` and resets every relevant
    module-level cache/singleton so `_generate_export_async` (and this
    test's own MinIO client) genuinely talk to THIS module's dedicated
    containers — same rationale as `tests/
    test_schedule_runs_progress_sse.py`'s `_point_at_real_redis` fixture.
    """

    monkeypatch.setenv("RPD_DATABASE_URL", postgres_url)
    monkeypatch.setenv("RPD_MINIO_ENDPOINT", minio_endpoint)
    monkeypatch.setenv("RPD_MINIO_ACCESS_KEY", _MINIO_ACCESS_KEY)
    monkeypatch.setenv("RPD_MINIO_SECRET_KEY", _MINIO_SECRET_KEY)
    monkeypatch.setenv("RPD_MINIO_SECURE", "false")
    get_minio_settings.cache_clear()
    minio_config_module._minio_client = None
    yield
    minio_config_module._minio_client = None
    get_minio_settings.cache_clear()


async def _seed_matrix_project(postgres_url: str):
    """Real, committed rows (not a rolled-back SAVEPOINT) — this task's own
    fresh engine (opened inside `_generate_export_async`, on ITS OWN event
    loop) must be able to see them from a completely separate connection.
    """

    engine = create_async_engine(postgres_url)
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        async with session_factory() as db:
            hub = Hub(name=HubName.RD_GREECE, lab_region=LabRegion.GREECE, is_oem=False)
            db.add(hub)
            await db.flush()

            user = User(
                email=f"export-worker-{uuid.uuid4().hex[:8]}@example.com",
                full_name="Export Worker Test User",
                oidc_subject=None,
                is_active=True,
                hub_scope_all=True,
            )
            db.add(user)
            await db.flush()

            project = Project(
                name="Worker Task Matrix Project",
                hub_id=hub.id,
                category=ProjectCategory.A,
                type=ProjectType.NM,
                status=ProjectStatus.IN_QUEUE,
                priority=ProjectPriority.P2,
                frozen=False,
                delay_weeks=0,
                carry_over=False,
                tcogs_eur=999.99,
            )
            db.add(project)
            await db.flush()

            score = PriorityScore(
                project_id=project.id,
                **dict.fromkeys(DIMENSION_FIELD_NAMES, 4),
                hard_gates=[],
            )
            db.add(score)
            await db.flush()

            job = ExportJob(
                surface=ExportSurface.MATRIX,
                format=ExportFormat.XLSX,
                requested_by_user_id=user.id,
                filters={"currency": None, "hub_id": None, "category": None},
                status=ExportJobStatus.QUEUED,
            )
            db.add(job)
            await db.commit()
            await db.refresh(job)
            await db.refresh(user)
            return job.id, user.id
    finally:
        await engine.dispose()


def _principal_payload(user_id: uuid.UUID) -> dict:
    return {
        "user_id": str(user_id),
        "email": "export-worker@example.com",
        "full_name": "Export Worker Test User",
        "oidc_subject": "test-subject",
        "roles": [RoleName.PORTFOLIO_MANAGER.value],
        "hub_scope_all": True,
        "hub_ids": [],
    }


async def test_generate_export_uploads_to_minio_and_completes(_point_at_real_infra, postgres_url):
    job_id, user_id = await _seed_matrix_project(postgres_url)

    result = await _generate_export_async(
        _FakeTask(),
        export_job_id=str(job_id),
        principal=_principal_payload(user_id),
        surface="matrix",
        format="xlsx",
        filters={"currency": None, "hub_id": None, "category": None},
    )
    assert result["status"] == "completed"
    assert result["row_count"] == 1

    # --- verify the ExportJob row transitioned correctly, via a fresh
    # connection (never the task's own, already-disposed engine) ------------
    engine = create_async_engine(postgres_url)
    try:
        session_factory = async_sessionmaker(engine, expire_on_commit=False)
        async with session_factory() as db:
            job = await db.get(ExportJob, job_id)
            assert job.status == ExportJobStatus.COMPLETED
            assert job.row_count == 1
            assert job.object_key == f"exports/{job_id}.xlsx"
            assert job.celery_task_id == "fake-worker-task-id"
            assert job.completed_at is not None

            # --- verify the object is genuinely retrievable from MinIO, with
            # the expected content (not just that a row/key was recorded) ---
            client = get_minio_client()
            settings = get_minio_settings()
            response = client.get_object(settings.exports_bucket, job.object_key)
            try:
                content = response.read()
            finally:
                response.close()
                response.release_conn()
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
            assert workbook.sheetnames == ["Matrix"]
            ws = workbook["Matrix"]
            rows = list(ws.iter_rows(values_only=True))
            headers = list(rows[0])
            data_row = dict(zip(headers, rows[1], strict=True))
            assert data_row["project_name"] == "Worker Task Matrix Project"
            assert float(data_row["tcogs_eur"]) == pytest.approx(999.99)

            # --- verify the real HTTP-shaped download endpoint (called
            # directly, not through ASGI — see module docstring) streams the
            # SAME bytes back to the job's own requester ---------------------
            principal = Principal(
                user_id=user_id,
                email="export-worker@example.com",
                full_name="Export Worker Test User",
                oidc_subject="test-subject",
                roles=frozenset({RoleName.PORTFOLIO_MANAGER}),
                hub_scope_all=True,
            )
            streaming_response = await download_export_job(job_id, db=db, current_user=principal)
            downloaded = b"".join([chunk async for chunk in streaming_response.body_iterator])
            assert downloaded == content
            assert (
                streaming_response.headers["content-disposition"]
                == 'attachment; filename="matrix-export.xlsx"'
            )
    finally:
        await engine.dispose()


async def test_generate_export_marks_job_failed_on_builder_error(
    _point_at_real_infra, postgres_url
):
    """`surface="not-a-real-surface"` (a task-call argument, independent of
    the `ExportJob.surface` DB column, which stays a valid enum value) makes
    `services.export_builder.build_export_sheets` raise `ValueError` —
    confirms the task's broad `except Exception` handler marks the row
    `FAILED` with a truncated `error_message`, rather than leaving it stuck
    in `RUNNING` forever (mirroring `workers/schedule_tasks.py`'s own
    documented "never leave a row stuck" contract).
    """

    engine = create_async_engine(postgres_url)
    try:
        session_factory = async_sessionmaker(engine, expire_on_commit=False)
        async with session_factory() as db:
            user = User(
                email=f"export-worker-fail-{uuid.uuid4().hex[:8]}@example.com",
                full_name="Export Worker Fail Test User",
                oidc_subject=None,
                is_active=True,
                hub_scope_all=True,
            )
            db.add(user)
            await db.flush()
            job = ExportJob(
                surface=ExportSurface.DASHBOARD,
                format=ExportFormat.CSV,
                requested_by_user_id=user.id,
                filters={},
                status=ExportJobStatus.QUEUED,
            )
            db.add(job)
            await db.commit()
            await db.refresh(job)
            job_id, user_id = job.id, user.id
    finally:
        await engine.dispose()

    with pytest.raises(ValueError, match="Unknown export surface"):
        await _generate_export_async(
            _FakeTask(),
            export_job_id=str(job_id),
            principal=_principal_payload(user_id),
            surface="not-a-real-surface",
            format="csv",
            filters={},
        )

    engine = create_async_engine(postgres_url)
    try:
        session_factory = async_sessionmaker(engine, expire_on_commit=False)
        async with session_factory() as db:
            job = await db.get(ExportJob, job_id)
            assert job.status == ExportJobStatus.FAILED
            assert job.error_message is not None
            assert "Unknown export surface" in job.error_message
            assert job.completed_at is not None
    finally:
        await engine.dispose()
