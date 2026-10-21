"""P25-T03 (ADR 0036, docs/API_CONTRACT_P25.md §5): the P1 capacity fill is
served everywhere a band is shown: the Priority board card, the Workspace
score section and `GET /exports/matrix`, from the same one-pass loader as
the Matrix (`services.priority_capacity_fill`).

Reuses P25-T01's hand-worked world (hub capacity 10.00 weeks; Base 80 / 6 wk,
First 65 / 3 wk promoted to 9.00, Zero 62 / 0 wk skipped, Second 60 / 2 wk
stopped at 11.00, Low 30 / 1 wk not reached).
"""

from __future__ import annotations

import csv
import io

import pytest
from sqlalchemy import event

from api.db import get_db
from api.main import app
from models.enums import HubName, LabRegion, RoleName
from services.export_builder import MATRIX_PROMOTED_BAND_HEADER
from tests.auth_helpers import clear_current_principal_override
from tests.factories import make_schedule_run
from tests.test_p25_capacity_fill import (
    _act_as,
    _client,
    _hand_worked_world,
    _hub10,
    _project,
)

FILL_FIELDS = ("suggested_band", "score_band", "capacity_fill", "capacity_fill_reason")


@pytest.fixture(autouse=True)
def _cleanup():
    yield
    clear_current_principal_override()
    app.dependency_overrides.pop(get_db, None)


def _board_cards(board: dict) -> dict[str, dict]:
    return {c["name"]: c for col in board["columns"] for c in col["cards"]}


def _export_rows(content: bytes) -> tuple[list[str], dict[str, dict[str, str]]]:
    reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
    rows = list(reader)
    return list(reader.fieldnames or []), {r["project_name"]: r for r in rows}


def _csv_fill(row: dict[str, str]) -> tuple[str | None, str | None, bool, str | None]:
    return (
        row[MATRIX_PROMOTED_BAND_HEADER] or None,
        row["score_band"] or None,
        row["capacity_fill"] == "True",
        row["capacity_fill_reason"] or None,
    )


async def test_promoted_project_is_the_same_on_every_surface(db_session):
    _hub, _run, p = await _hand_worked_world(db_session)
    p["first"].external_code = "S-4012"
    await db_session.flush()
    await _act_as(db_session, RoleName.PORTFOLIO_MANAGER)
    async with _client(db_session) as client:
        matrix = {r["project_name"]: r for r in (await client.get("/priorities")).json()}
        board = _board_cards((await client.get("/priority-board")).json())
        export = await client.get("/exports/matrix", params={"format": "csv"})
        workspaces = {
            name: (await client.get(f"/projects/{proj.id}/workspace")).json()["priority_score"]
            for name, proj in (("First", p["first"]), ("Second", p["second"]), ("Base", p["base"]))
        }
    assert export.status_code == 200
    headers, exported = _export_rows(export.content)

    for name in ("First", "Second", "Zero", "Base", "Low"):
        expected = tuple(matrix[name][f] for f in FILL_FIELDS)
        assert tuple(board[name][f] for f in FILL_FIELDS) == expected, name
        assert _csv_fill(exported[name]) == expected, name
        if name in workspaces:
            assert tuple(workspaces[name][f] for f in FILL_FIELDS) == expected, name

    # The promoted project itself, spelled out.
    assert tuple(board["First"][f] for f in FILL_FIELDS) == (
        "P1",
        "P2",
        True,
        "P1 (capacity fill): hub P1 design load 9.00 of 10.00 weeks with this project",
    )
    # The board keeps the stored score band in `band` (= `score_band`).
    assert board["First"]["band"] == "P2" == board["First"]["score_band"]
    assert board["First"]["external_code"] == "S-4012"
    assert board["Second"]["external_code"] is None
    # Stopped candidate: reason, no promotion.
    assert workspaces["Second"]["capacity_fill"] is False
    assert workspaces["Second"]["capacity_fill_reason"] == (
        "Not promoted: hub P1 design load would be 11.00 of 10.00 weeks"
    )

    # Export columns: the promoted band is labelled as such; the plain
    # `suggested_band` header is gone; the fill columns follow it.
    assert "suggested_band" not in headers
    i = headers.index(MATRIX_PROMOTED_BAND_HEADER)
    assert headers[i : i + 4] == [
        MATRIX_PROMOTED_BAND_HEADER,
        "score_band",
        "capacity_fill",
        "capacity_fill_reason",
    ]
    # Nested fields stay out of the spreadsheet.
    assert "dimension_points" not in headers and "dimension_scores" not in headers


async def test_export_xlsx_carries_the_fill_and_neutralises_formulas(db_session):
    """The new text column goes through the same formula guard; a formula-like
    project name is still neutralised in the Matrix export."""

    import openpyxl

    hub, run, p = await _hand_worked_world(db_session)
    await _project(db_session, hub, run, "=HYPERLINK(1)", 20, 1)
    await _act_as(db_session, RoleName.PORTFOLIO_MANAGER)
    async with _client(db_session) as client:
        resp = await client.get("/exports/matrix", params={"format": "xlsx"})
    ws = openpyxl.load_workbook(io.BytesIO(resp.content))["Matrix"]
    rows = list(ws.iter_rows(values_only=True))
    headers = list(rows[0])
    by_name = {
        r[headers.index("project_name")]: dict(zip(headers, r, strict=True)) for r in rows[1:]
    }
    first = by_name["First"]
    assert (first[MATRIX_PROMOTED_BAND_HEADER], first["score_band"], first["capacity_fill"]) == (
        "P1",
        "P2",
        True,
    )
    assert first["capacity_fill_reason"].startswith("P1 (capacity fill)")
    assert "'=HYPERLINK(1)" in by_name


async def test_board_totals_are_unchanged_by_the_fill(db_session):
    """Columns stay keyed by the committed priority: the promoted project
    remains in its committed column and counts/weeks are the committed ones,
    with or without an active run (with no run the fill is off)."""

    await _hand_worked_world(db_session)
    await _act_as(db_session, RoleName.PORTFOLIO_MANAGER)
    async with _client(db_session) as client:
        board = (await client.get("/priority-board")).json()
    totals = {
        c["priority"]: (c["count"], c["design_weeks"], c["lab_weeks"]) for c in board["columns"]
    }
    # Every hand-worked project is committed P3; remaining design 6+3+0+2+1.
    assert totals == {
        "P1": (0, 0.0, 0.0),
        "P2": (0, 0.0, 0.0),
        "P3": (5, 12.0, 0.0),
        "P4": (0, 0.0, 0.0),
        "Q": (0, 0.0, 0.0),
    }
    p3 = next(c for c in board["columns"] if c["priority"] == "P3")
    assert [c["name"] for c in p3["cards"] if c["capacity_fill"]] == ["First"]


async def test_board_and_workspace_with_no_active_run(db_session):
    _hub, _run, p = await _hand_worked_world(db_session, active=False)
    await _act_as(db_session, RoleName.PORTFOLIO_MANAGER)
    async with _client(db_session) as client:
        board = _board_cards((await client.get("/priority-board")).json())
        ws = (await client.get(f"/projects/{p['first'].id}/workspace")).json()["priority_score"]
    for card in board.values():
        assert card["capacity_fill"] is False and card["capacity_fill_reason"] is None
        assert card["suggested_band"] == card["score_band"] == card["band"]
    assert (ws["suggested_band"], ws["score_band"], ws["capacity_fill"]) == ("P2", "P2", False)


async def test_hub_planner_board_is_scoped_and_fill_matches_admin(db_session):
    greece, run, _p = await _hand_worked_world(db_session)
    romania = await _hub10(db_session, HubName.PD_ROMANIA, LabRegion.ROMANIA)
    await _project(db_session, romania, run, "RO cand", 60, 4)
    await _act_as(db_session, RoleName.PORTFOLIO_MANAGER)
    async with _client(db_session) as client:
        admin = _board_cards((await client.get("/priority-board")).json())
    assert admin["RO cand"]["capacity_fill"] is True

    await _act_as(db_session, RoleName.HUB_PLANNER, hub_ids=[greece.id])
    async with _client(db_session) as client:
        mine = _board_cards((await client.get("/priority-board")).json())
        other = (await client.get("/priority-board", params={"hub_id": str(romania.id)})).json()
    assert {c["hub"] for c in mine.values()} == {"R&D-Greece"}
    assert all(col["cards"] == [] for col in other["columns"])
    for name, card in mine.items():
        assert {f: card[f] for f in FILL_FIELDS} == {f: admin[name][f] for f in FILL_FIELDS}


async def test_workspace_unscored_project_has_no_score_section(db_session):
    from tests.factories import make_project

    hub, _run, _p = await _hand_worked_world(db_session)
    bare = await make_project(db_session, hub, name="Unscored")
    await _act_as(db_session, RoleName.PORTFOLIO_MANAGER)
    async with _client(db_session) as client:
        resp = await client.get(f"/projects/{bare.id}/workspace")
    assert resp.status_code == 200 and resp.json()["priority_score"] is None


def _counting(db_session):
    statements: list[str] = []

    def _count(_conn, _cursor, statement, *_args):
        statements.append(statement)

    return statements, _count


async def test_board_and_workspace_query_counts_are_constant(db_session):
    """No N+1: 3 vs 30 projects in the hub give the same statement count for
    the board and for one project's Workspace."""

    hub = await _hub10(db_session)
    run = await make_schedule_run(db_session, is_active=True)
    anchor = await _project(db_session, hub, run, "Anchor", 64, 1)
    for i in range(2):
        await _project(db_session, hub, run, f"S{i}", 60 + i, 1)
    await _act_as(db_session, RoleName.PORTFOLIO_MANAGER)

    statements, _count = _counting(db_session)
    engine = db_session.bind.sync_engine

    async def _measure(client) -> tuple[int, int]:
        statements.clear()
        event.listen(engine, "before_cursor_execute", _count)
        try:
            assert (await client.get("/priority-board")).status_code == 200
            n_board = len(statements)
            statements.clear()
            assert (await client.get(f"/projects/{anchor.id}/workspace")).status_code == 200
            n_ws = len(statements)
        finally:
            event.remove(engine, "before_cursor_execute", _count)
        return n_board, n_ws

    async with _client(db_session) as client:
        small = await _measure(client)
        for i in range(27):
            await _project(
                db_session, hub, run, f"L{i}", 40 + i % 30, 1 + i % 3, external_code=f"C-{i}"
            )
        large = await _measure(client)
        board = _board_cards((await client.get("/priority-board")).json())
    assert small == large
    assert len(board) == 30
